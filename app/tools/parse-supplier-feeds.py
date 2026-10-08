#!/usr/bin/env python3
"""Offline, provenance-preserving supplier feed extraction. No shop/PIM writes."""
from pathlib import Path
from collections import Counter
from decimal import Decimal, InvalidOperation
import argparse, hashlib, json, os, re, zipfile
from defusedxml import ElementTree as ET
import openpyxl

LIMIT = 256 * 1024 * 1024
IDENTITY_PARAMS = {'Артикул', 'Код', 'Код товару', 'SKU'}
COLOR_PARAMS = {'Колір', 'Цвет', 'Color'}
SIZE_PARAMS = {'Розмір', 'Розміри', 'Размер', 'Размеры', 'Size'}
BARCODE_PARAMS = {'Штрихкод', 'Штрих-код', 'Barcode', 'EAN', 'GTIN'}

def text(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return 'true' if value else 'false'
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None

def quantity(value):
    raw = text(value)
    if raw is None or not re.fullmatch(r'\d+(?:\.\d+)?', raw):
        return None
    try:
        return format(Decimal(raw).normalize(), 'f')
    except InvalidOperation:
        return None

def values(params, names):
    return list(dict.fromkeys(p['value'] for p in params if p['name'] in names and p['value']))

def local(tag):
    return tag.rsplit('}', 1)[-1]

def xml_feed(path):
    offers, categories, generated = [], [], None
    # YML commonly declares shops.dtd. Never resolve external DTDs or entities.
    for event, element in ET.iterparse(path, events=('start', 'end'), forbid_dtd=False,
                                       forbid_entities=True, forbid_external=True):
        tag = local(element.tag)
        if event == 'start' and tag == 'yml_catalog':
            generated = element.attrib.get('date')
        if event != 'end':
            continue
        if tag == 'category':
            categories.append({'id': element.attrib.get('id'), 'parent_id': element.attrib.get('parentId'), 'name': text(element.text)})
        if tag != 'offer':
            continue
        fields, params = {}, []
        for child in element:
            name, value = local(child.tag), text(child.text)
            if name == 'param':
                params.append({'name': text(child.attrib.get('name')), 'value': value, 'attributes': dict(child.attrib)})
            else:
                fields.setdefault(name, []).append(value)
        first = lambda name: next((v for v in fields.get(name, []) if v is not None), None)
        identities = [{'field': '@id', 'value': text(element.attrib.get('id'))}]
        for name in ['vendorCode', 'article']:
            identities += [{'field': name, 'value': v} for v in fields.get(name, []) if v]
        identities += [{'field': 'param:' + p['name'], 'value': p['value']} for p in params if p['name'] in IDENTITY_PARAMS and p['value']]
        qty_values = [(name, first(name)) for name in ['quantity_in_stock', 'stock_quantity', 'quantity', 'stock'] if first(name) is not None]
        parsed_qty = {quantity(v) for _, v in qty_values}
        qty = next(iter(parsed_qty)) if len(parsed_qty) == 1 else None
        offers.append({'offer_id': text(element.attrib.get('id')), 'supplier_sku': first('vendorCode') or text(element.attrib.get('id')),
                       'identity_fields': identities, 'group_id': text(element.attrib.get('group_id')),
                       'group_source': '@group_id' if element.attrib.get('group_id') else None,
                       'name': first('name_ua') or first('name'), 'brand': first('vendor'),
                       'category_id': first('categoryId'), 'url': first('url'), 'photos': [v for v in fields.get('picture', []) if v],
                       'barcodes': list(dict.fromkeys([v for v in fields.get('barcode', []) if v] + values(params, BARCODE_PARAMS))),
                       'color_values': values(params, COLOR_PARAMS), 'size_values': values(params, SIZE_PARAMS),
                       'params': params, 'fields': fields, 'attributes': dict(element.attrib),
                       'availability_raw': element.attrib.get('available'), 'in_stock_raw': element.attrib.get('in_stock'),
                       'quantity': qty, 'quantity_fields': [{'field': k, 'raw': v} for k, v in qty_values],
                       'quantity_verified_current': False})
        element.clear()
    return offers, categories, generated

def xlsx_feed(path, sku_field):
    with zipfile.ZipFile(path) as archive:
        files = archive.infolist()
        if len(files) > 10000 or sum(f.file_size for f in files) > LIMIT or any(f.filename.endswith('vbaProject.bin') for f in files):
            raise ValueError('Workbook archive exceeds limits or contains VBA')
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    offers, categories = [], []
    try:
        if 'Export Products Sheet' not in book.sheetnames:
            raise ValueError('Expected Prom.ua Export Products Sheet')
        rows = book['Export Products Sheet'].iter_rows(values_only=True)
        headers = list(next(rows))
        for row_number, row in enumerate(rows, 2):
            if not any(value is not None for value in row):
                continue
            fields = {}
            for i, value in enumerate(row):
                if i < len(headers) and headers[i]:
                    fields.setdefault(headers[i], []).append(value)
            first = lambda key: next((text(v) for v in fields.get(key, []) if text(v) is not None), None)
            params = []
            for i, header in enumerate(headers):
                if header == 'Назва_Характеристики' and i + 2 < len(row) and text(row[i]):
                    params.append({'name': text(row[i]), 'value': text(row[i+2]), 'unit': text(row[i+1])})
            identities = [{'field': name, 'value': first(name)} for name in ['Унікальний_ідентифікатор', 'Ідентифікатор_товару', 'Код_товару', 'Номер_пристрою_(MPN)'] if first(name)]
            group = first('ID_групи_різновидів')
            photo_text = first('Посилання_зображення') or ''
            offers.append({'offer_id': first('Унікальний_ідентифікатор') or first('Ідентифікатор_товару'), 'supplier_sku': first(sku_field),
                           'source_row_number': row_number, 'identity_fields': identities, 'group_id': group,
                           'group_source': 'ID_групи_різновидів' if group else None,
                           'name': first('Назва_позиції_укр') or first('Назва_позиції'), 'brand': first('Виробник'),
                           'category_id': first('Ідентифікатор_підрозділу'),
                           'url': None, 'photos': [p.strip() for p in photo_text.split(',') if p.strip()],
                           'barcodes': [first('Код_маркування_(GTIN)')] if first('Код_маркування_(GTIN)') else [],
                           'color_values': values(params, COLOR_PARAMS), 'size_values': values(params, SIZE_PARAMS),
                           'params': params, 'fields': fields, 'attributes': {}, 'availability_raw': first('Наявність'),
                           'quantity': quantity(first('Кількість')),
                           'quantity_fields': [{'field': 'Кількість', 'raw': first('Кількість')}] if first('Кількість') is not None else [],
                           'quantity_verified_current': False})
        if 'Export Groups Sheet' in book.sheetnames:
            rows = book['Export Groups Sheet'].iter_rows(values_only=True)
            headers = list(next(rows))
            for row in rows:
                if any(v is not None for v in row):
                    categories.append(dict(zip(headers, row)))
    finally:
        book.close()
    # Workbook category/group columns do not imply MODEL groups or export date.
    return offers, categories, None

def parse_feed(binding):
    path = Path(binding['file']).resolve()
    if path.stat().st_size > LIMIT:
        raise ValueError('Input exceeds 256 MiB')
    if path.suffix.lower() == '.xml':
        offers, categories, generated = xml_feed(path)
    elif path.suffix.lower() == '.xlsx':
        offers, categories, generated = xlsx_feed(path, binding['sku_field'])
    else:
        raise ValueError('Unsupported supplier file')
    ids = Counter(o['offer_id'] for o in offers if o['offer_id'])
    groups = Counter(o['group_id'] for o in offers if o['group_id'])
    errors = [{'code': 'DUPLICATE_OFFER_ID', 'id': key} for key, n in ids.items() if n > 1]
    errors += [{'code': 'MISSING_OFFER_ID', 'row': i+1} for i, o in enumerate(offers) if not o['offer_id']]
    return {'supplier_id': binding['supplier_id'], 'supplier_key': binding['supplier_key'],
            'source_sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'source_bytes': path.stat().st_size,
            'generated_at_raw': generated, 'freshness_verified': False, 'offers': offers, 'categories': categories, 'errors': errors,
            'summary': {'offers': len(offers), 'explicit_model_groups': len(groups), 'grouped_offers': sum(groups.values()),
                        'offers_without_group': sum(not o['group_id'] for o in offers),
                        'quantity_observations': sum(o['quantity'] is not None for o in offers),
                        'conflicting_colors': sum(len(o['color_values']) > 1 for o in offers),
                        'errors': len(errors)}}

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--bindings', required=True)
    parser.add_argument('--output-dir', required=True)
    args = parser.parse_args()
    bindings = json.loads(Path(args.bindings).read_text())
    if not isinstance(bindings, list) or not bindings:
        raise ValueError('Explicit supplier bindings list required')
    ids = [b['supplier_id'] for b in bindings]
    if len(ids) != len(set(ids)):
        raise ValueError('Duplicate supplier binding')
    output = Path(args.output_dir).resolve()
    root = Path(__file__).resolve().parents[1]
    if output == root or root in output.parents:
        raise ValueError('Output must be outside the application/web root')
    feeds = [parse_feed(binding) for binding in bindings]
    os.umask(0o077)
    output.mkdir(mode=0o700)
    data = {'format': 'rubizh.supplier-sources.v2', 'contains_customer_data': False, 'feeds': feeds}
    with (output / 'feeds.private.json').open('x') as file:
        json.dump(data, file, ensure_ascii=False, separators=(',', ':'), default=str)
    summary = {'feeds': [{'supplier_key': f['supplier_key'], 'source_sha256': f['source_sha256'], 'generated_at_raw': f['generated_at_raw'], **f['summary']} for f in feeds],
               'errors': sum(len(f['errors']) for f in feeds), 'production_modified': False, 'freshness_verified': False}
    with (output / 'summary.json').open('x') as file:
        json.dump(summary, file, ensure_ascii=False, indent=2)
    print(json.dumps(summary, ensure_ascii=False))

if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        raise SystemExit('Supplier extraction stopped: ' + str(error))
